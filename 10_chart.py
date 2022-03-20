from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference

wb = load_workbook("test.xlsx")
ws = wb.active

""" bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
bar_chart = BarChart()
bar_chart.add_data(bar_value)

ws.add_chart(bar_chart, "F1") """

line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data=True)
line_chart.title = "성적표"
line_chart.style = 20
line_chart.y_axis.title="점수"
line_chart.x_axis.title="번호"

ws.add_chart(line_chart, "F1")

wb.save("new.xlsx")
wb.close()